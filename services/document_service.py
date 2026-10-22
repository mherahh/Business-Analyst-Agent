"""Functions for creating Word documents."""
from docx import Document
from config.settings import OUTPUT_DIR

def save_doc(filename, title, sections):
    """Create a Word document from a title and heading/line pairs."""
    doc = Document()
    doc.add_heading(title, 0)
    for heading, lines in sections:
        doc.add_heading(heading, 1)
        for line in lines:
            doc.add_paragraph(str(line), style="List Bullet")
    path = f"{OUTPUT_DIR}/{filename}"
    doc.save(path)
    return path
